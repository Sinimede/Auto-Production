
import unittest
import os
import openpyxl
from src.utils.excel_writer import generate_excel, find_header_row

class TestExcelWriter(unittest.TestCase):
    def setUp(self):
        self.test_dir = "tests/utils/tmp"
        if not os.path.exists(self.test_dir):
            os.makedirs(self.test_dir)
        
        self.template_path = os.path.join(self.test_dir, "template.xlsx")
        self.output_path = os.path.join(self.test_dir, "output.xlsx")
        
        # Create a mock template
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.cell(row=3, column=2, value="Pos")
        ws.cell(row=3, column=3, value="Qt")
        wb.save(self.template_path)

    def tearDown(self):
        if os.path.exists(self.test_dir):
            import shutil, stat, time, gc
            gc.collect()
            def _on_error(func, path, exc_info):
                os.chmod(path, stat.S_IWRITE)
                time.sleep(0.05)
                func(path)
            shutil.rmtree(self.test_dir, onerror=_on_error)

    def test_find_header_row(self):
        wb = openpyxl.load_workbook(self.template_path)
        ws = wb.active
        row = find_header_row(ws)
        wb.close()
        self.assertEqual(row, 3)

    def test_generate_excel_dynamic_header(self):
        rows = [{"qty": 2, "part_number": "P123"}]
        column_order = ["qty", "part_number"]
        
        header_row = generate_excel(self.template_path, rows, self.output_path, column_order)
        
        self.assertEqual(header_row, 3)
        
        wb = openpyxl.load_workbook(self.output_path)
        ws = wb.active
        val_qty = ws.cell(row=4, column=2).value
        val_pn = ws.cell(row=4, column=3).value
        wb.close()
        self.assertEqual(val_qty, 2)
        self.assertEqual(val_pn, "P123")

if __name__ == '__main__':
    unittest.main()
