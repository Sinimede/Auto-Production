# tools/fornecedores/generate_fornecedores.py
"""
Bootstrap script — run once to generate Templates/Fornecedores.xlsx.
Re-running overwrites the file; manual corrections in the Excel are lost.
"""
import os
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(__file__))
from categorias import CATEGORIAS, VALID_CATEGORIES, DEFAULT_CATEGORY

# ── Colours ────────────────────────────────────────────────────────────────
HEADER_FILL  = PatternFill("solid", fgColor="D9D9D9")
CAT_FILLS = {
    "Mecânico":   PatternFill("solid", fgColor="CCE5FF"),
    "Elétrico":   PatternFill("solid", fgColor="FFFACC"),
    "Pneumático": PatternFill("solid", fgColor="CCFFCC"),
}

# ── Helpers ─────────────────────────────────────────────────────────────────

def derive_code(name: str) -> str:
    """Return uppercase 3-letter (or shorter) prefix from supplier name."""
    return name[:3].upper()


def load_suppliers(txt_path: str) -> list:
    """Read supplier names from txt file; strip whitespace, skip blanks, deduplicate (keep first)."""
    seen = set()
    result = []
    with open(txt_path, encoding="utf-8", errors="replace") as f:
        for line in f:
            name = line.strip()
            if not name:
                continue
            if name not in seen:
                seen.add(name)
                result.append(name)
    return result


def build_rows(suppliers: list) -> list:
    """Return sorted list of (nome, código, categoria) tuples."""
    rows = []
    for name in suppliers:
        code = derive_code(name)
        cat  = CATEGORIAS.get(name, DEFAULT_CATEGORY)
        rows.append((name, code, cat))
    rows.sort(key=lambda r: r[0].lower())
    return rows


def write_excel(rows: list, out_path: str) -> None:
    """Write the formatted Excel file."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fornecedores"

    # Header
    headers = ["Nome", "Código", "Categoria"]
    ws.append(headers)
    for col, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")

    # Data validation dropdown on column C
    dv = DataValidation(
        type="list",
        formula1='"Mecânico,Elétrico,Pneumático"',
        allow_blank=False,
        showDropDown=False,
    )
    ws.add_data_validation(dv)

    # Data rows
    for i, (nome, codigo, categoria) in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=nome)
        ws.cell(row=i, column=2, value=codigo)
        cat_cell = ws.cell(row=i, column=3, value=categoria)
        fill = CAT_FILLS.get(categoria, CAT_FILLS["Mecânico"])
        for col in range(1, 4):
            ws.cell(row=i, column=col).fill = fill
        dv.add(cat_cell)

    # Column widths
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 15

    wb.save(out_path)
    print(f"Saved {len(rows)} suppliers -> {out_path}")


def main():
    base     = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    txt_path = os.path.join(base, "Templates", "Fornecedores.txt")
    out_path = os.path.join(base, "Templates", "Fornecedores.xlsx")

    if not os.path.exists(txt_path):
        print(f"ERROR: source file not found: {txt_path}")
        sys.exit(1)

    suppliers = load_suppliers(txt_path)
    rows      = build_rows(suppliers)
    write_excel(rows, out_path)


if __name__ == "__main__":
    main()
