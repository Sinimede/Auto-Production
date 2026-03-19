"""
excel_writer.py — Excel generation from templates.

Copies a template xlsx, fills part data rows, and saves to output_path.
Never modifies the original template. Overwrites output_path if it exists.
"""

import os
import re
import shutil
import openpyxl

from paths import get_templates_dir

# Column orders per process.
# Each list maps to openpyxl columns starting at B (index 2).
# Keys must match the dict returned by solidworks.get_part_data().

LASER_COLUMNS = [
    "qty", "part_number", "Description", "Corte_Fabrico",
    "Simetria", "Material", "TratSuperficial", "espessura",
]

ROUTER_COLUMNS = [
    "qty", "part_number", "Description", "Corte_Fabrico",
    "espessura", "Material", "Simetria",
]

CNC_COLUMNS = [
    "qty", "part_number", "Description", "Corte_Fabrico",
    "Material", "TratSuperficial", "Simetria",
]

TORNO_COLUMNS = [
    "qty", "part_number", "Description", "Corte_Fabrico",
    "Material", "TratSuperficial", "Simetria",
]

PROTECOES_COLUMNS = [
    "qty", "part_number", "Description", "Material",
]

# Maps process keyword -> (template filename, output filename, column order, data_start_row)
PROCESS_CONFIG = {
    "laser":     ("Laser_template.xlsx",     "Laser.xlsx",     LASER_COLUMNS,     4),
    "router":    ("Router_template.xlsx",    "Router.xlsx",    ROUTER_COLUMNS,    3),
    "cnc":       ("CNC_template.xlsx",       "CNC.xlsx",       CNC_COLUMNS,       4),
    "torno":     ("Torno_template.xlsx",     "Torno.xlsx",     TORNO_COLUMNS,     4),
    "protecoes": ("Protecoes_template.xlsx", "Protecoes.xlsx", PROTECOES_COLUMNS, 4),
}


def generate_excel(template_path: str, rows: list, output_path: str,
                   column_order: list, data_start_row: int = 4) -> None:
    """
    Copy template_path to output_path, then write part data rows.

    Layout:
      - Rows before data_start_row (title/header) are preserved from the template.
      - Data starts at data_start_row, column B (openpyxl index 2).
      - column_order[0] -> col B (index 2), column_order[1] -> col C (index 3), etc.
      - espessura None -> written as empty string "".
      - Overwrites output_path silently if it already exists.

    Raises: OSError on copy failure; openpyxl errors on workbook save.
    """
    shutil.copy2(template_path, output_path)
    wb = openpyxl.load_workbook(output_path)
    ws = wb["Folha1"]

    for row_offset, row_dict in enumerate(rows):
        row_idx = data_start_row + row_offset
        for col_offset, key in enumerate(column_order):
            value = row_dict.get(key, "")
            if value is None:
                value = ""
            ws.cell(row=row_idx, column=2 + col_offset).value = value

    # Remove blank template rows that follow the data.
    # If there are more data rows than the template had pre-allocated, ws.max_row
    # will already reflect the extended content and first_blank > ws.max_row — no
    # rows are deleted, which is correct.
    first_blank = data_start_row + len(rows)
    if first_blank <= ws.max_row:
        ws.delete_rows(first_blank, ws.max_row - first_blank + 1)

    # Update Excel Table refs to match the actual data range.
    # Templates define a Tabela1 Table object with a fixed ref (e.g. B3:I43).
    # After deleting rows, the Table ref must be updated — otherwise Excel
    # regenerates blank rows from the stale ref when the file is opened.
    header_row = data_start_row - 1
    last_row = header_row if not rows else data_start_row + len(rows) - 1
    for tbl in ws.tables.values():
        m = re.match(r'([A-Z]+)\d+:([A-Z]+)\d+', tbl.ref)
        if m:
            tbl.ref = f"{m.group(1)}{header_row}:{m.group(2)}{last_row}"

    wb.save(output_path)
