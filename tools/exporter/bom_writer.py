"""
bom_writer.py — BOM (Lista de Material) Excel generation.

Copies the macro template (.xlsm), writes a plain workbook (.xlsx).
VBA is intentionally stripped via keep_vba=False — no macros needed in output.
No SolidWorks COM dependency — pure openpyxl + stdlib.
"""
import os
import re
import shutil
import zipfile
import logging

import openpyxl
from openpyxl.styles import Alignment

from paths import get_templates_dir

# ---------------------------------------------------------------------------
# Brand classification — backed by Templates/Fornecedores.xlsx
# ---------------------------------------------------------------------------

_CAT_MAP = {
    "Mecânico":   "mecanico",
    "Elétrico":   "eletrico",
    "Pneumático": "pneumatico",
}


def normalize(s: str) -> str:
    return s.lower().replace(" ", "")


def load_fornecedores_lookup() -> dict:
    """
    Load Templates/Fornecedores.xlsx and return {normalize(nome): category}.
    category is one of "mecanico", "eletrico", "pneumatico".
    Returns {} and logs a warning if the file is missing.
    """
    path = os.path.join(get_templates_dir(), "Fornecedores.xlsx")
    if not os.path.isfile(path):
        logging.warning("AVISO: Fornecedores.xlsx não encontrado — classificação por defeito: Mecânico")
        return {}
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    result = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        nome = row[0]
        cat_pt = row[2] if len(row) > 2 else None
        if not nome or not cat_pt:
            continue
        internal = _CAT_MAP.get(str(cat_pt).strip())
        if internal:
            result[normalize(str(nome))] = internal
    wb.close()
    return result


FORNECEDORES_LOOKUP: dict = load_fornecedores_lookup()
ALL_KNOWN_BRANDS: frozenset = frozenset(FORNECEDORES_LOOKUP.keys())


def is_known_brand(corte_fabrico) -> bool:
    if isinstance(corte_fabrico, tuple):
        corte_fabrico = corte_fabrico[0] if corte_fabrico else ""
    corte_fabrico = corte_fabrico or ""
    val = normalize(corte_fabrico)
    if val in FORNECEDORES_LOOKUP:
        return True
    return any(n and n in val for n in FORNECEDORES_LOOKUP)


def classify_commercial(corte_fabrico) -> str:
    """
    Return "pneumatico", "eletrico", or "mecanico".
    Exact match first; then substring scan (alphabetical, first wins).
    Unknown / empty falls back to "mecanico".
    """
    if isinstance(corte_fabrico, tuple):
        corte_fabrico = corte_fabrico[0] if corte_fabrico else ""
    corte_fabrico = corte_fabrico or ""
    val = normalize(corte_fabrico)
    if val in FORNECEDORES_LOOKUP:
        return FORNECEDORES_LOOKUP[val]
    for norm_name in sorted(FORNECEDORES_LOOKUP):
        if norm_name and norm_name in val:
            return FORNECEDORES_LOOKUP[norm_name]
    return "mecanico"

# ---------------------------------------------------------------------------
# Column definitions
# ---------------------------------------------------------------------------

# None entries → write "" for that column
PRODUCAO_COLUMNS = [
    "qty",             # col 1 (A) — Quant.
    "part_number",     # col 2 (B) — Número Interno
    None,              # col 3 (C) — Rev. (always empty — dept. convention)
    "Description",     # col 4 (D) — Descrição
    "Corte_Fabrico",   # col 5 (E) — Corte/Fabrico
    "Simetria",        # col 6 (F) — Simetria
    "Material",        # col 7 (G) — Material
    "TratSuperficial", # col 8 (H) — Tratamento Superficial
    "A_Partir_de",     # col 9 (I) — A partir de
]

COMERCIAL_COLUMNS = [
    "qty",           # col 1 (A) — Quant.
    "part_number",   # col 2 (B) — Número Interno
    None,            # col 3 (C) — Coluna1 (Excel auto-name, no business meaning, always empty)
    "Description",   # col 4 (D) — Referência/nº Proposta
    "Corte_Fabrico", # col 5 (E) — Marca/Fabricante
    None,            # col 6 (F) — Descrição (always empty — dept. convention)
]

# ---------------------------------------------------------------------------
# Excel generation
# ---------------------------------------------------------------------------

def generate_bom(rows_producao: list, rows_mecanico: list, rows_eletrico: list,
                 rows_pneumatico: list, output_path: str) -> None:
    """
    Write Lista de materiais.xlsx from the macro template.
    All row list args are list[dict]. Empty lists are allowed.
    Raises FileNotFoundError if template is missing.
    """
    template_path = os.path.join(get_templates_dir(), "Lista-de-Material_template.xlsm")
    shutil.copy2(template_path, output_path)   # raises FileNotFoundError if missing
    wb = openpyxl.load_workbook(output_path, keep_vba=False)

    sheet_config = [
        ("Produção",            rows_producao,   PRODUCAO_COLUMNS),
        ("Material Mecânico",   rows_mecanico,   COMERCIAL_COLUMNS),
        ("Material Elétrico",   rows_eletrico,   COMERCIAL_COLUMNS),
        ("Material Pneumático", rows_pneumatico, COMERCIAL_COLUMNS),
    ]

    for sheet_name, rows, col_keys in sheet_config:
        ws = wb[sheet_name]
        data_start_row = 3
        header_row = 2   # invariant for all 4 sheets

        for row_offset, row_dict in enumerate(rows):
            row_idx = data_start_row + row_offset
            for col_offset, key in enumerate(col_keys):
                value = row_dict.get(key, "") if key else ""
                cell = ws.cell(row=row_idx, column=1 + col_offset)
                cell.value = value or ""
                cell.alignment = Alignment(horizontal="center")

        # Remove trailing blank template rows
        first_blank = data_start_row + len(rows)
        if first_blank <= ws.max_row:
            ws.delete_rows(first_blank, ws.max_row - first_blank + 1)

        # Update Excel Table refs to match actual data range.
        # Column letters are preserved from the template (e.g. "A2:I1098" → "A2:I7").
        # Column bounds are NOT hardcoded — extracted via regex from existing ref.
        last_row = header_row if not rows else data_start_row + len(rows) - 1
        for tbl in ws.tables.values():
            m = re.match(r'([A-Z]+)\d+:([A-Z]+)\d+', tbl.ref)
            if m:
                tbl.ref = f"{m.group(1)}{header_row}:{m.group(2)}{last_row}"
        # If a sheet has no Excel Table the loop is a no-op.

    wb.save(output_path)

    # Post-save XML patch: openpyxl writes empty string cells as self-closing
    # inlineStr elements (e.g. <c ... t="inlineStr" />) which load back as None.
    # Replace them with <is><t/></is> content so they round-trip as "".
    _patch_empty_inlinestr(output_path)


def _patch_empty_inlinestr(xlsx_path: str) -> None:
    """
    Rewrite xlsx_path in-place, replacing self-closing inlineStr cells with
    proper empty-string XML so openpyxl reads them back as '' instead of None.
    """
    tmp_path = xlsx_path + ".tmp"
    try:
        with zipfile.ZipFile(xlsx_path, "r") as zin:
            with zipfile.ZipFile(tmp_path, "w", zipfile.ZIP_DEFLATED) as zout:
                for item in zin.infolist():
                    data = zin.read(item.filename)
                    if item.filename.startswith("xl/worksheets/sheet") and item.filename.endswith(".xml"):
                        text = data.decode("utf-8")
                        # Self-closing inlineStr: t="inlineStr" /> or t="inlineStr"/>
                        text = re.sub(
                            r't="inlineStr"\s*/>',
                            't="inlineStr"><is><t/></is></c>',
                            text
                        )
                        data = text.encode("utf-8")
                    zout.writestr(item, data)
        os.replace(tmp_path, xlsx_path)
    except Exception:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)
        raise
