"""
Unit tests for excel_writer.generate_excel.
Uses the real template files from Templates/ — no COM required.

Template layouts (confirmed from actual files):
  Laser:  header row 3, data starts row 4
  Router: header row 2, data starts row 3
  CNC:    header row 3, data starts row 4
  Torno:  header row 3, data starts row 4
"""
import sys
import os
import shutil
import tempfile
import unittest
import openpyxl

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))

# Templates are at repo root / Templates/
TEMPLATES_DIR = os.path.join(
    os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(
        os.path.abspath(__file__))))),
    "Templates"
)

LASER_TEMPLATE  = os.path.join(TEMPLATES_DIR, "Laser_template.xlsx")
ROUTER_TEMPLATE = os.path.join(TEMPLATES_DIR, "Router_template.xlsx")
CNC_TEMPLATE    = os.path.join(TEMPLATES_DIR, "CNC_template.xlsx")
TORNO_TEMPLATE  = os.path.join(TEMPLATES_DIR, "Torno_template.xlsx")


class TestGenerateExcel(unittest.TestCase):

    def setUp(self):
        self.tmp = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.tmp)

    # --- Laser (data_start_row=4) ---

    def test_laser_data_written_to_correct_columns(self):
        from excel_writer import generate_excel, LASER_COLUMNS
        output = os.path.join(self.tmp, "Laser.xlsx")
        rows = [{
            "qty": 2, "part_number": "12624.100.001",
            "Description": "Chapa Base", "Corte_Fabrico": "LASER",
            "Simetria": "", "Material": "S235JR",
            "TratSuperficial": "Pintar RAL 9002", "espessura": 3.0,
        }]
        generate_excel(LASER_TEMPLATE, rows, output, LASER_COLUMNS, data_start_row=4)

        ws = openpyxl.load_workbook(output)["Folha1"]
        self.assertEqual(ws.cell(row=4, column=2).value, 2)
        self.assertEqual(ws.cell(row=4, column=3).value, "12624.100.001")
        self.assertEqual(ws.cell(row=4, column=4).value, "Chapa Base")
        self.assertEqual(ws.cell(row=4, column=5).value, "LASER")
        self.assertIn(ws.cell(row=4, column=6).value, ("", None))
        self.assertEqual(ws.cell(row=4, column=7).value, "S235JR")
        self.assertEqual(ws.cell(row=4, column=8).value, "Pintar RAL 9002")
        self.assertEqual(ws.cell(row=4, column=9).value, 3.0)

    def test_header_row_preserved(self):
        from excel_writer import generate_excel, LASER_COLUMNS
        output = os.path.join(self.tmp, "Laser.xlsx")
        generate_excel(LASER_TEMPLATE, [], output, LASER_COLUMNS, data_start_row=4)
        ws = openpyxl.load_workbook(output)["Folha1"]
        self.assertIsNotNone(ws.cell(row=3, column=2).value)

    def test_original_template_not_modified(self):
        from excel_writer import generate_excel, LASER_COLUMNS
        output = os.path.join(self.tmp, "Laser.xlsx")
        wb_before = openpyxl.load_workbook(LASER_TEMPLATE)
        header_before = wb_before["Folha1"].cell(row=3, column=2).value

        rows = [{"qty": 99, "part_number": "X", "Description": "X",
                 "Corte_Fabrico": "X", "Simetria": "X", "Material": "X",
                 "TratSuperficial": "X", "espessura": 1.0}]
        generate_excel(LASER_TEMPLATE, rows, output, LASER_COLUMNS, data_start_row=4)

        wb_after = openpyxl.load_workbook(LASER_TEMPLATE)
        self.assertEqual(wb_after["Folha1"].cell(row=3, column=2).value, header_before)

    def test_espessura_none_written_as_empty_string(self):
        from excel_writer import generate_excel, LASER_COLUMNS
        output = os.path.join(self.tmp, "Laser.xlsx")
        rows = [{"qty": 1, "part_number": "X", "Description": "", "Corte_Fabrico": "",
                 "Simetria": "", "Material": "", "TratSuperficial": "", "espessura": None}]
        generate_excel(LASER_TEMPLATE, rows, output, LASER_COLUMNS, data_start_row=4)
        ws = openpyxl.load_workbook(output)["Folha1"]
        self.assertIn(ws.cell(row=4, column=9).value, ("", None))

    def test_multiple_rows_written_in_order(self):
        from excel_writer import generate_excel, LASER_COLUMNS
        output = os.path.join(self.tmp, "Laser.xlsx")
        rows = [
            {"qty": 1, "part_number": "P1", "Description": "", "Corte_Fabrico": "",
             "Simetria": "", "Material": "", "TratSuperficial": "", "espessura": None},
            {"qty": 2, "part_number": "P2", "Description": "", "Corte_Fabrico": "",
             "Simetria": "", "Material": "", "TratSuperficial": "", "espessura": None},
        ]
        generate_excel(LASER_TEMPLATE, rows, output, LASER_COLUMNS, data_start_row=4)
        ws = openpyxl.load_workbook(output)["Folha1"]
        self.assertEqual(ws.cell(row=4, column=3).value, "P1")
        self.assertEqual(ws.cell(row=5, column=3).value, "P2")

    def test_overwrites_existing_file(self):
        from excel_writer import generate_excel, LASER_COLUMNS
        output = os.path.join(self.tmp, "Laser.xlsx")
        rows_v1 = [{"qty": 1, "part_number": "OLD", "Description": "", "Corte_Fabrico": "",
                    "Simetria": "", "Material": "", "TratSuperficial": "", "espessura": None}]
        generate_excel(LASER_TEMPLATE, rows_v1, output, LASER_COLUMNS, data_start_row=4)
        rows_v2 = [{"qty": 2, "part_number": "NEW", "Description": "", "Corte_Fabrico": "",
                    "Simetria": "", "Material": "", "TratSuperficial": "", "espessura": None}]
        generate_excel(LASER_TEMPLATE, rows_v2, output, LASER_COLUMNS, data_start_row=4)
        ws = openpyxl.load_workbook(output)["Folha1"]
        self.assertEqual(ws.cell(row=4, column=3).value, "NEW")

    # --- Router (data_start_row=3) ---

    def test_router_columns_espessura_at_col_f(self):
        from excel_writer import generate_excel, ROUTER_COLUMNS
        output = os.path.join(self.tmp, "Router.xlsx")
        rows = [{"qty": 1, "part_number": "P", "Description": "D",
                 "Corte_Fabrico": "Router", "espessura": 15.0,
                 "Material": "Aluminio", "Simetria": "P_mirror"}]
        generate_excel(ROUTER_TEMPLATE, rows, output, ROUTER_COLUMNS, data_start_row=3)
        ws = openpyxl.load_workbook(output)["Folha1"]
        self.assertEqual(ws.cell(row=3, column=6).value, 15.0)
        self.assertEqual(ws.cell(row=3, column=8).value, "P_mirror")

    # --- CNC (data_start_row=4) ---

    def test_cnc_columns_no_espessura(self):
        from excel_writer import generate_excel, CNC_COLUMNS
        output = os.path.join(self.tmp, "CNC.xlsx")
        rows = [{"qty": 2, "part_number": "C1", "Description": "Block",
                 "Corte_Fabrico": "CNC", "Material": "AL5083",
                 "TratSuperficial": "Anodizar", "Simetria": ""}]
        generate_excel(CNC_TEMPLATE, rows, output, CNC_COLUMNS, data_start_row=4)
        ws = openpyxl.load_workbook(output)["Folha1"]
        self.assertEqual(ws.cell(row=4, column=2).value, 2)
        self.assertEqual(ws.cell(row=4, column=6).value, "AL5083")
        self.assertEqual(ws.cell(row=4, column=7).value, "Anodizar")

    # --- Torno ---

    def test_torno_same_column_order_as_cnc(self):
        from excel_writer import generate_excel, TORNO_COLUMNS, CNC_COLUMNS
        self.assertEqual(TORNO_COLUMNS, CNC_COLUMNS)

    # --- Blank row trimming ---

    def _make_row(self, part_number):
        return {"qty": 1, "part_number": part_number, "Description": "",
                "Corte_Fabrico": "", "Simetria": "", "Material": "",
                "TratSuperficial": "", "espessura": None}

    def test_output_has_no_blank_rows_after_data(self):
        """Table must end at last data row — no trailing blank rows from template."""
        from excel_writer import generate_excel, LASER_COLUMNS
        output = os.path.join(self.tmp, "Laser.xlsx")
        rows = [self._make_row("P1"), self._make_row("P2"), self._make_row("P3")]
        generate_excel(LASER_TEMPLATE, rows, output, LASER_COLUMNS, data_start_row=4)
        ws = openpyxl.load_workbook(output)["Folha1"]
        # 3 data rows starting at row 4 → last row with data = row 6
        self.assertEqual(ws.max_row, 6)

    def test_output_header_plus_one_row(self):
        """Single part: table = header row + 1 data row."""
        from excel_writer import generate_excel, LASER_COLUMNS
        output = os.path.join(self.tmp, "Laser_one.xlsx")
        rows = [self._make_row("SOLO")]
        generate_excel(LASER_TEMPLATE, rows, output, LASER_COLUMNS, data_start_row=4)
        ws = openpyxl.load_workbook(output)["Folha1"]
        self.assertEqual(ws.max_row, 4)  # header row 3 + 1 data row 4

    def test_output_exceeds_template_row_count(self):
        """More parts than template pre-allocated rows — all rows written, none deleted."""
        from excel_writer import generate_excel, ROUTER_COLUMNS
        output = os.path.join(self.tmp, "Router_big.xlsx")
        # Router template has max_row=10, data_start_row=3 → 8 pre-allocated data rows
        # Write 15 parts — output must have 15 data rows
        rows = [{"qty": 1, "part_number": f"P{i}", "Description": "",
                 "Corte_Fabrico": "", "espessura": None, "Material": "", "Simetria": ""}
                for i in range(15)]
        generate_excel(ROUTER_TEMPLATE, rows, output, ROUTER_COLUMNS, data_start_row=3)
        ws = openpyxl.load_workbook(output)["Folha1"]
        self.assertEqual(ws.max_row, 17)  # header row 2 + 15 data rows
        self.assertEqual(ws.cell(row=17, column=3).value, "P14")

    def test_zero_rows_keeps_only_header(self):
        """No parts — only header row remains, all pre-allocated template rows deleted."""
        from excel_writer import generate_excel, LASER_COLUMNS
        output = os.path.join(self.tmp, "Laser_empty.xlsx")
        generate_excel(LASER_TEMPLATE, [], output, LASER_COLUMNS, data_start_row=4)
        ws = openpyxl.load_workbook(output)["Folha1"]
        self.assertEqual(ws.max_row, 3)  # rows 1-2 empty + row 3 header

    # --- Table ref updates ---

    def test_table_ref_ends_at_last_data_row(self):
        """Table ref must be updated so Excel does not regenerate blank rows."""
        import re
        from excel_writer import generate_excel, LASER_COLUMNS
        output = os.path.join(self.tmp, "Laser_tbl.xlsx")
        rows = [self._make_row("P1"), self._make_row("P2"), self._make_row("P3")]
        generate_excel(LASER_TEMPLATE, rows, output, LASER_COLUMNS, data_start_row=4)
        ws = openpyxl.load_workbook(output)["Folha1"]
        for tbl in ws.tables.values():
            m = re.search(r':([A-Z]+)(\d+)$', tbl.ref)
            self.assertEqual(int(m.group(2)), 6,
                             f"Table ref {tbl.ref} should end at row 6 for 3 data rows")

    def test_table_ref_zero_rows_ends_at_header(self):
        """Empty table ref should end at header row (no data rows)."""
        import re
        from excel_writer import generate_excel, LASER_COLUMNS
        output = os.path.join(self.tmp, "Laser_tbl_empty.xlsx")
        generate_excel(LASER_TEMPLATE, [], output, LASER_COLUMNS, data_start_row=4)
        ws = openpyxl.load_workbook(output)["Folha1"]
        for tbl in ws.tables.values():
            m = re.search(r':([A-Z]+)(\d+)$', tbl.ref)
            self.assertEqual(int(m.group(2)), 3,
                             f"Table ref {tbl.ref} should end at header row 3 when no data")

    def test_table_ref_router_ends_correctly(self):
        """Router: header row 2, 2 data rows → table should end at row 4."""
        import re
        from excel_writer import generate_excel, ROUTER_COLUMNS
        output = os.path.join(self.tmp, "Router_tbl.xlsx")
        rows = [{"qty": 1, "part_number": "P1", "Description": "", "Corte_Fabrico": "",
                 "espessura": None, "Material": "", "Simetria": ""},
                {"qty": 1, "part_number": "P2", "Description": "", "Corte_Fabrico": "",
                 "espessura": None, "Material": "", "Simetria": ""}]
        generate_excel(ROUTER_TEMPLATE, rows, output, ROUTER_COLUMNS, data_start_row=3)
        ws = openpyxl.load_workbook(output)["Folha1"]
        for tbl in ws.tables.values():
            m = re.search(r':([A-Z]+)(\d+)$', tbl.ref)
            self.assertEqual(int(m.group(2)), 4,
                             f"Table ref {tbl.ref} should end at row 4")


PROTECOES_TEMPLATE = os.path.join(TEMPLATES_DIR, "Protecoes_template.xlsx")


class TestProtecoesConfig(unittest.TestCase):
    """PROTECOES_COLUMNS and PROCESS_CONFIG['protecoes'] are correctly defined."""

    def test_protecoes_columns_defined(self):
        from excel_writer import PROTECOES_COLUMNS
        self.assertEqual(PROTECOES_COLUMNS,
                         ["qty", "part_number", "Description", "Material"])

    def test_process_config_has_protecoes(self):
        from excel_writer import PROCESS_CONFIG
        self.assertIn("protecoes", PROCESS_CONFIG)

    def test_process_config_protecoes_values(self):
        from excel_writer import PROCESS_CONFIG, PROTECOES_COLUMNS
        tmpl, out, cols, start_row = PROCESS_CONFIG["protecoes"]
        self.assertEqual(tmpl, "Protecoes_template.xlsx")
        self.assertEqual(out, "Protecoes.xlsx")
        self.assertEqual(cols, PROTECOES_COLUMNS)
        self.assertEqual(start_row, 4)


class TestGenerateExcelProtecoes(unittest.TestCase):

    def setUp(self):
        self.tmp = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.tmp)

    def test_protecoes_data_written_to_correct_columns(self):
        from excel_writer import generate_excel, PROTECOES_COLUMNS
        output = os.path.join(self.tmp, "Protecoes.xlsx")
        rows = [{
            "qty": 3, "part_number": "18026.200.001",
            "Description": "Tampa Proteção", "Material": "S235JR",
        }]
        generate_excel(PROTECOES_TEMPLATE, rows, output,
                       PROTECOES_COLUMNS, data_start_row=4)

        ws = openpyxl.load_workbook(output)["Folha1"]
        self.assertEqual(ws.cell(row=4, column=2).value, 3)
        self.assertEqual(ws.cell(row=4, column=3).value, "18026.200.001")
        self.assertEqual(ws.cell(row=4, column=4).value, "Tampa Proteção")
        self.assertEqual(ws.cell(row=4, column=5).value, "S235JR")

    def test_header_row_preserved(self):
        from excel_writer import generate_excel, PROTECOES_COLUMNS
        output = os.path.join(self.tmp, "Protecoes_hdr.xlsx")
        generate_excel(PROTECOES_TEMPLATE, [], output,
                       PROTECOES_COLUMNS, data_start_row=4)
        ws = openpyxl.load_workbook(output)["Folha1"]
        # Row 3 col B should be the header (QTY.)
        self.assertIsNotNone(ws.cell(row=3, column=2).value)

    def test_original_template_not_modified(self):
        from excel_writer import generate_excel, PROTECOES_COLUMNS
        output = os.path.join(self.tmp, "Protecoes_orig.xlsx")
        hdr_before = openpyxl.load_workbook(
            PROTECOES_TEMPLATE)["Folha1"].cell(row=3, column=2).value
        rows = [{"qty": 1, "part_number": "X", "Description": "X", "Material": "X"}]
        generate_excel(PROTECOES_TEMPLATE, rows, output,
                       PROTECOES_COLUMNS, data_start_row=4)
        hdr_after = openpyxl.load_workbook(
            PROTECOES_TEMPLATE)["Folha1"].cell(row=3, column=2).value
        self.assertEqual(hdr_before, hdr_after)


if __name__ == "__main__":
    unittest.main()
