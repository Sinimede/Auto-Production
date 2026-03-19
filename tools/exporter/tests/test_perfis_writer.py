"""
Unit tests for perfis_writer.py.
Uses the real Perfis-de-Alumínio_template.xlsx — SolidWorks not needed.
"""
import sys
import os
import tempfile
import unittest
import openpyxl

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))

TEMPLATES_DIR = os.path.join(
    os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(
        os.path.abspath(__file__))))),
    "Templates"
)
TEMPLATE_PATH = os.path.join(TEMPLATES_DIR, "Perfis-de-Alumínio_template.xlsx")


class TestGeneratePerfis(unittest.TestCase):

    def setUp(self):
        self.tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        self.tmp.close()
        self.out_path = self.tmp.name

    def tearDown(self):
        try:
            os.unlink(self.out_path)
        except OSError:
            pass

    def _load(self):
        return openpyxl.load_workbook(self.out_path)

    def test_creates_output_file(self):
        from perfis_writer import generate_perfis
        generate_perfis([], self.out_path)
        self.assertTrue(os.path.isfile(self.out_path))

    def test_column_a_not_written(self):
        from perfis_writer import generate_perfis
        rows = [{"qty": 2, "description": "Profile A", "length_mm": 100.0}]
        generate_perfis(rows, self.out_path)
        wb = self._load()
        ws = wb.active
        self.assertIsNone(ws["A2"].value)
        self.assertIsNone(ws["A3"].value)

    def test_writes_data_row_columns(self):
        from perfis_writer import generate_perfis
        rows = [{"qty": 2, "description": "Aluminium profile 40x40", "length_mm": 350.0}]
        generate_perfis(rows, self.out_path)
        wb = self._load()
        ws = wb.active
        self.assertEqual(ws["B3"].value, 2)          # qty
        self.assertEqual(ws["C3"].value, "Aluminium profile 40x40")
        self.assertEqual(ws["D3"].value, 350.0)      # length_mm

    def test_length_rounded_to_2dp(self):
        from perfis_writer import generate_perfis
        rows = [{"qty": 1, "description": "Profile A", "length_mm": 566.4285}]
        generate_perfis(rows, self.out_path)
        wb = self._load()
        ws = wb.active
        self.assertAlmostEqual(ws["D3"].value, 566.43, places=2)

    def test_table_ref_updated_when_overflow(self):
        """More than 22 data rows → table ref must grow beyond D24."""
        from perfis_writer import generate_perfis
        rows = [{"qty": 1, "description": f"Profile {i}", "length_mm": float(i * 10)}
                for i in range(1, 30)]  # 29 rows — beyond template's 22
        generate_perfis(rows, self.out_path)
        wb = self._load()
        ws = wb.active
        for tbl in ws.tables.values():
            # last data row = 3 + 29 - 1 = 31
            self.assertIn("31", tbl.ref)

    def test_empty_rows_table_ref_collapses_to_header(self):
        """Empty input → table ref shrinks to just the header row."""
        from perfis_writer import generate_perfis
        generate_perfis([], self.out_path)
        wb = self._load()
        ws = wb.active
        for tbl in ws.tables.values():
            # last_row = HEADER_ROW = 2 when no data
            self.assertTrue(tbl.ref.endswith("2"), msg=f"ref={tbl.ref}")


if __name__ == "__main__":
    unittest.main()
