"""
Unit tests for bom_writer.py.
Uses the real template file from Templates/ — SolidWorks does not need to be running.
"""
import sys
import os
import shutil
import tempfile
import unittest
import openpyxl

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))

TEMPLATES_DIR = os.path.join(
    os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(
        os.path.abspath(__file__))))),
    "Templates"
)
BOM_TEMPLATE = os.path.join(TEMPLATES_DIR, "Lista-de-Material_template.xlsm")


# ---------------------------------------------------------------------------
# classify_commercial
# ---------------------------------------------------------------------------

class TestClassifyCommercial(unittest.TestCase):

    def test_festo_uppercase_is_pneumatico(self):
        from bom_writer import classify_commercial
        self.assertEqual(classify_commercial("FESTO"), "pneumatico")

    def test_festo_lowercase_is_pneumatico(self):
        from bom_writer import classify_commercial
        self.assertEqual(classify_commercial("festo"), "pneumatico")

    def test_festo_mixed_case_is_pneumatico(self):
        from bom_writer import classify_commercial
        self.assertEqual(classify_commercial("Festo"), "pneumatico")

    def test_smc_is_pneumatico(self):
        from bom_writer import classify_commercial
        self.assertEqual(classify_commercial("SMC"), "pneumatico")

    def test_siemens_is_eletrico(self):
        from bom_writer import classify_commercial
        self.assertEqual(classify_commercial("Siemens"), "eletrico")

    def test_balluf_is_eletrico(self):
        from bom_writer import classify_commercial
        self.assertEqual(classify_commercial("Balluf"), "eletrico")

    def test_skf_is_mecanico(self):
        from bom_writer import classify_commercial
        self.assertEqual(classify_commercial("SKF"), "mecanico")

    def test_bosch_rexroth_is_mecanico(self):
        from bom_writer import classify_commercial
        self.assertEqual(classify_commercial("Bosch Rexroth"), "mecanico")

    def test_empty_string_is_mecanico(self):
        from bom_writer import classify_commercial
        self.assertEqual(classify_commercial(""), "mecanico")

    def test_completely_unknown_brand_is_mecanico(self):
        from bom_writer import classify_commercial
        self.assertEqual(classify_commercial("XYZ Unknown Corp"), "mecanico")


# ---------------------------------------------------------------------------
# ALL_KNOWN_BRANDS — unknown brand detection
# ---------------------------------------------------------------------------

class TestAllKnownBrands(unittest.TestCase):

    def test_xyz_not_in_all_known_brands(self):
        from bom_writer import ALL_KNOWN_BRANDS
        val = "xyz"
        self.assertFalse(any(b in val for b in ALL_KNOWN_BRANDS))

    def test_festo_in_all_known_brands(self):
        from bom_writer import ALL_KNOWN_BRANDS
        val = "festo"
        self.assertTrue(any(b in val for b in ALL_KNOWN_BRANDS))

    def test_skf_in_all_known_brands(self):
        from bom_writer import ALL_KNOWN_BRANDS
        val = "skf"
        self.assertTrue(any(b in val for b in ALL_KNOWN_BRANDS))


# ---------------------------------------------------------------------------
# generate_bom — uses real template
# ---------------------------------------------------------------------------

class TestGenerateBom(unittest.TestCase):

    def setUp(self):
        if not os.path.isfile(BOM_TEMPLATE):
            self.skipTest(f"Template not found: {BOM_TEMPLATE}")
        self.tmp = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.tmp)

    def _output(self, name="Lista de materiais.xlsx"):
        return os.path.join(self.tmp, name)

    def test_output_file_is_openable_xlsx(self):
        from bom_writer import generate_bom
        out = self._output()
        generate_bom([], [], [], [], out)
        # openpyxl must open without keep_vba (VBA stripped)
        wb = openpyxl.load_workbook(out)
        wb.close()

    def test_producao_columns_written_correctly(self):
        from bom_writer import generate_bom
        out = self._output()
        rows_prod = [{
            "qty": 3,
            "part_number": "18026.100.001",
            "Description": "Chapa Base",
            "Corte_Fabrico": "LASER",
            "Simetria": "S",
            "Material": "S235JR",
            "TratSuperficial": "Pintar RAL 9002",
            "A_Partir_de": "Chapa 3mm",
        }]
        generate_bom(rows_prod, [], [], [], out)
        wb = openpyxl.load_workbook(out)
        ws = wb["Produção"]
        # Data starts at row 3
        self.assertEqual(ws.cell(3, 1).value, 3)                  # qty
        self.assertEqual(ws.cell(3, 2).value, "18026.100.001")    # part_number
        self.assertEqual(ws.cell(3, 3).value, "")                 # Rev (always empty)
        self.assertEqual(ws.cell(3, 4).value, "Chapa Base")       # Description
        self.assertEqual(ws.cell(3, 5).value, "LASER")            # Corte_Fabrico
        self.assertEqual(ws.cell(3, 6).value, "S")                # Simetria
        self.assertEqual(ws.cell(3, 7).value, "S235JR")           # Material
        self.assertEqual(ws.cell(3, 8).value, "Pintar RAL 9002")  # TratSuperficial
        self.assertEqual(ws.cell(3, 9).value, "Chapa 3mm")        # A_Partir_de
        wb.close()

    def test_comercial_mecanico_columns_written_correctly(self):
        from bom_writer import generate_bom
        out = self._output()
        rows_mec = [{
            "qty": 2,
            "part_number": "18026.800.001",
            "Description": "REF-123",
            "Corte_Fabrico": "SKF",
        }]
        generate_bom([], rows_mec, [], [], out)
        wb = openpyxl.load_workbook(out)
        ws = wb["Material Mecânico"]
        self.assertEqual(ws.cell(3, 1).value, 2)                # qty
        self.assertEqual(ws.cell(3, 2).value, "18026.800.001")  # part_number
        self.assertEqual(ws.cell(3, 3).value, "")               # Coluna1 (always empty)
        self.assertEqual(ws.cell(3, 4).value, "REF-123")        # Description/Referência
        self.assertEqual(ws.cell(3, 5).value, "SKF")            # Corte_Fabrico/Marca
        self.assertEqual(ws.cell(3, 6).value, "")               # Descrição (always empty)
        wb.close()

    def test_none_value_written_as_empty_string(self):
        from bom_writer import generate_bom
        out = self._output()
        rows_prod = [{
            "qty": 1,
            "part_number": "18026.100.002",
            "Description": None,
            "Corte_Fabrico": "CNC",
            "Simetria": None,
            "Material": "S235JR",
            "TratSuperficial": None,
            "A_Partir_de": None,
        }]
        generate_bom(rows_prod, [], [], [], out)
        wb = openpyxl.load_workbook(out)
        ws = wb["Produção"]
        self.assertEqual(ws.cell(3, 4).value, "")  # Description → ""
        self.assertEqual(ws.cell(3, 6).value, "")  # Simetria → ""
        self.assertEqual(ws.cell(3, 8).value, "")  # TratSuperficial → ""
        self.assertEqual(ws.cell(3, 9).value, "")  # A_Partir_de → ""
        wb.close()

    def test_empty_producao_removes_template_blanks(self):
        from bom_writer import generate_bom
        out = self._output()
        generate_bom([], [], [], [], out)
        wb = openpyxl.load_workbook(out)
        ws = wb["Produção"]
        # No data rows → all rows from row 3 deleted → max_row = 2
        self.assertEqual(ws.max_row, 2)
        wb.close()

    def test_table_ref_header_only_when_no_data(self):
        from bom_writer import generate_bom
        out = self._output()
        generate_bom([], [], [], [], out)
        wb = openpyxl.load_workbook(out)
        ws = wb["Produção"]
        for tbl in ws.tables.values():
            # header_row=2, no data → last_row=2 → ref ends with "2"
            self.assertTrue(
                tbl.ref.endswith("2"),
                f"Expected ref ending in '2', got {tbl.ref}"
            )
        wb.close()

    def test_table_ref_updated_with_two_data_rows(self):
        from bom_writer import generate_bom
        out = self._output()
        row_template = {
            "qty": 1, "part_number": "18026.100.001", "Description": "A",
            "Corte_Fabrico": "LASER", "Simetria": "", "Material": "S235JR",
            "TratSuperficial": "", "A_Partir_de": "",
        }
        rows_prod = [dict(row_template), dict(row_template, part_number="18026.100.002")]
        generate_bom(rows_prod, [], [], [], out)
        wb = openpyxl.load_workbook(out)
        ws = wb["Produção"]
        # 2 data rows → last_row = data_start_row(3) + 2 - 1 = 4
        for tbl in ws.tables.values():
            self.assertTrue(
                tbl.ref.endswith("4"),
                f"Expected ref ending in '4', got {tbl.ref}"
            )
        wb.close()

    def test_missing_key_in_row_dict_written_as_empty(self):
        from bom_writer import generate_bom
        out = self._output()
        # Minimal dict — optional keys missing
        rows_prod = [{"qty": 1, "part_number": "18026.100.001"}]
        generate_bom(rows_prod, [], [], [], out)
        wb = openpyxl.load_workbook(out)
        ws = wb["Produção"]
        # Description (col D = 4) should be "" when key absent
        self.assertEqual(ws.cell(3, 4).value, "")
        wb.close()

    def test_template_missing_raises_file_not_found(self):
        from bom_writer import generate_bom
        import unittest.mock as mock
        out = self._output()
        with mock.patch("bom_writer.get_templates_dir", return_value="/nonexistent/path"):
            with self.assertRaises(FileNotFoundError):
                generate_bom([], [], [], [], out)


if __name__ == "__main__":
    unittest.main()
