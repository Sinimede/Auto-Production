
import os
import re
import shutil
import openpyxl
from src.utils.path_utils import normalize_path

# Column orders per process.
LASER_COLUMNS = [
    "qty", "part_number", "Description", "Corte_Fabrico",
    "Simetria", "Material", "TratSuperficial", "espessura",
]

ROUTER_COLUMNS = [
    "qty", "part_number", "Description", "Corte_Fabrico",
    "espessura", "Material", "Simetria",
]

CNC_COLUMNS = [
    "qty", "part_number", "Description", "Corte_Fabrico",
    "Material", "TratSuperficial", "Simetria",
]

TORNO_COLUMNS = [
    "qty", "part_number", "Description", "Corte_Fabrico",
    "Material", "TratSuperficial", "Simetria",
]

PROTECOES_COLUMNS = [
    "qty", "part_number", "Description", "Material",
]

PROCESS_CONFIG = {
    "laser":     ("Laser_template.xlsx",     "Laser.xlsx",     LASER_COLUMNS,     4),
    "router":    ("Router_template.xlsx",    "Router.xlsx",    ROUTER_COLUMNS,    3),
    "cnc":       ("CNC_template.xlsx",       "CNC.xlsx",       CNC_COLUMNS,       4),
    "torno":     ("Torno_template.xlsx",     "Torno.xlsx",     TORNO_COLUMNS,     4),
    "protecoes": ("Protecoes_template.xlsx", "Protecoes.xlsx", PROTECOES_COLUMNS, 4),
}

def find_header_row(sheet, max_search_rows=10):
    """
    Scans the sheet to find the header row by looking for 'Pos' or 'Qt'.
    Returns the 1-based row index or None if not found.
    """
    for r in range(1, max_search_rows + 1):
        # Column B (2) is usually where "Pos" or "Qt" is located
        cell_val = str(sheet.cell(row=r, column=2).value or "")
        if "Pos" in cell_val or "Qt" in cell_val:
            return r
    return None

def generate_excel(template_path: str, rows: list, output_path: str,
                   column_order: list, data_start_row: int = None) -> int:
    """
    Copy template_path to output_path, then write part data rows.
    Returns the header row index detected or used.
    """
    template_path = normalize_path(template_path)
    output_path = normalize_path(output_path)
    
    shutil.copy2(template_path, output_path)
    wb = openpyxl.load_workbook(output_path)
    ws = wb.active # Usually "Folha1"
    
    # Detect header row if not provided via data_start_row
    detected_header_row = find_header_row(ws)
    
    if data_start_row is None:
        if detected_header_row:
            data_start_row = detected_header_row + 1
        else:
            data_start_row = 4 # Default fallback
    
    effective_header_row = data_start_row - 1

    # Write data
    for row_offset, row_dict in enumerate(rows):
        row_idx = data_start_row + row_offset
        for col_offset, key in enumerate(column_order):
            value = row_dict.get(key, "")
            if value is None:
                value = ""
            # Data starts at Column B (2)
            ws.cell(row=row_idx, column=2 + col_offset).value = value

    # Remove blank template rows that follow the data
    first_blank = data_start_row + len(rows)
    if first_blank <= ws.max_row:
        ws.delete_rows(first_blank, ws.max_row - first_blank + 1)

    # Update Excel Table refs
    last_row = effective_header_row if not rows else data_start_row + len(rows) - 1
    for tbl in ws.tables.values():
        m = re.match(r'([A-Z]+)\d+:([A-Z]+)\d+', tbl.ref)
        if m:
            # We assume the table starts at the header row
            tbl.ref = f"{m.group(1)}{effective_header_row}:{m.group(2)}{last_row}"

    wb.save(output_path)
    return effective_header_row
