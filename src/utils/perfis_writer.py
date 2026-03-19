"""
perfis_writer.py — Perfis de Alumínio Excel generation.

Copies the Perfis-de-Alumínio_template.xlsx, writes weldment cut list rows.
No SolidWorks COM dependency — pure openpyxl + stdlib.

Template structure:
  Sheet "Folha1", Excel Table "Tabela1" ref B2:D24
  Row 2 (header): B2=QTY., C2=DESCRIPTION, D2=LENGTH
  Data rows start at row 3.
"""
import os
import re
import shutil

import openpyxl

from src.utils.path_utils import get_templates_dir

TEMPLATE_NAME  = "Perfis-de-Alumínio_template.xlsx"
HEADER_ROW     = 2
DATA_START_ROW = 3


def generate_perfis(rows: list, output_path: str) -> None:
    """
    Write Perfis de Alumínio Excel from the template.

    Args:
        rows:        List of dicts {"qty": int, "description": str, "length_mm": float}.
                     Must be pre-sorted (caller is responsible for sort order).
        output_path: Destination .xlsx path (overwritten if it exists).
    """
    template_path = os.path.join(get_templates_dir(), TEMPLATE_NAME)
    shutil.copy2(template_path, output_path)

    wb = openpyxl.load_workbook(output_path)
    ws = wb.active

    # Write data rows
    for i, row in enumerate(rows):
        r = DATA_START_ROW + i
        ws.cell(row=r, column=2).value = row["qty"]                     # QTY
        ws.cell(row=r, column=3).value = row["description"]             # DESCRIPTION
        ws.cell(row=r, column=4).value = round(row["length_mm"], 2)     # LENGTH

    # Update Excel Table ref to cover actual data range (same pattern as bom_writer.py)
    last_row = HEADER_ROW if not rows else DATA_START_ROW + len(rows) - 1
    for tbl in ws.tables.values():
        m = re.match(r'([A-Z]+)\d+:([A-Z]+)\d+', tbl.ref)
        if m:
            tbl.ref = f"{m.group(1)}{HEADER_ROW}:{m.group(2)}{last_row}"

    wb.save(output_path)
